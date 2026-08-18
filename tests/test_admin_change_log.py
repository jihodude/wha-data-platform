"""
test_admin_change_log.py — Change Log sheet in admin_inputs.xlsx.

Pins the contract for `record_change()` and the "Change Log" sheet:
every admin input mutation appends an audit row INSIDE the same workbook,
under a dedicated sheet that survives save/reload cycles.

Spec recap:
  - Sheet name: "Change Log"
  - Headers: Timestamp | User | Sheet | Cell | Previous Value | New Value
    (Reason retired 2026-08-13 — it only ever held a one-time migration
     note naming a person unrelated to later edits.)
  - record_change(workbook, sheet, cell, prev_value, new_value, user, reason="")
  - Idempotent for the same (timestamp_second, user, sheet, cell, new_value):
    duplicate calls within the same UTC second are dedup'd.
"""

from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from src.parsers.admin_inputs import (
    CHANGE_LOG_HEADERS,
    CHANGE_LOG_SHEET,
    ensure_change_log,
    record_change,
)


def _new_workbook_with_some_sheets():
    """Tiny fixture: a fresh workbook mimicking the admin xlsx (no Change Log yet)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Goals"
    wb.create_sheet("Member Goals")
    wb.create_sheet("Retention Goal")
    return wb


# ---------------------------------------------------------------------------
# ensure_change_log
# ---------------------------------------------------------------------------

def test_ensure_change_log_creates_sheet_with_headers():
    wb = _new_workbook_with_some_sheets()
    assert CHANGE_LOG_SHEET not in wb.sheetnames

    ensure_change_log(wb)

    assert CHANGE_LOG_SHEET in wb.sheetnames
    ws = wb[CHANGE_LOG_SHEET]
    actual_headers = [ws.cell(row=1, column=c + 1).value for c in range(len(CHANGE_LOG_HEADERS))]
    assert actual_headers == list(CHANGE_LOG_HEADERS)


def test_ensure_change_log_is_idempotent():
    wb = _new_workbook_with_some_sheets()
    ensure_change_log(wb)
    ensure_change_log(wb)  # second call should not duplicate or reset
    ws = wb[CHANGE_LOG_SHEET]
    # Only one header row, no data rows yet.
    assert ws.max_row == 1


# ---------------------------------------------------------------------------
# record_change — single call
# ---------------------------------------------------------------------------

def test_record_change_appends_one_row():
    wb = _new_workbook_with_some_sheets()
    record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah", reason="Q3 reforecast")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.max_row == 2  # header + one data row

    row = [ws.cell(row=2, column=c + 1).value for c in range(len(CHANGE_LOG_HEADERS))]
    timestamp, user, sheet, cell, prev_val, new_val = row

    # Timestamp must be ISO-8601 UTC parseable.
    assert isinstance(timestamp, str)
    parsed = datetime.fromisoformat(timestamp)
    assert parsed.tzinfo is not None  # tz-aware

    assert user == "hannah"
    assert sheet == "Sales Goals"
    assert cell == "B5"
    assert prev_val == 100
    assert new_val == 200
    assert len(row) == 6, "the retired Reason column must not come back"


def test_no_reason_column_is_written():
    wb = _new_workbook_with_some_sheets()
    record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
    ws = wb[CHANGE_LOG_SHEET]
    assert ws.cell(row=1, column=7).value is None
    assert ws.cell(row=2, column=7).value is None


def test_record_change_auto_creates_change_log_sheet():
    """Calling record_change on a workbook without the sheet should add it."""
    wb = _new_workbook_with_some_sheets()
    assert CHANGE_LOG_SHEET not in wb.sheetnames
    record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
    assert CHANGE_LOG_SHEET in wb.sheetnames


def test_record_change_handles_none_prev_value():
    """A blank cell becoming filled in is a valid edit."""
    wb = _new_workbook_with_some_sheets()
    record_change(wb, "Sales Goals", "C5", None, 500, user="jaima")
    ws = wb[CHANGE_LOG_SHEET]
    # openpyxl writes None as an empty cell; we accept either None or "".
    prev = ws.cell(row=2, column=5).value
    assert prev in (None, "")
    assert ws.cell(row=2, column=6).value == 500


# ---------------------------------------------------------------------------
# Idempotency / dedup
# ---------------------------------------------------------------------------

def test_record_change_dedups_duplicate_within_same_second():
    """
    Two identical record_change calls within the same UTC second collapse
    to a single row. (Same (timestamp_second, user, sheet, cell, new_value).)
    """
    wb = _new_workbook_with_some_sheets()
    fixed = datetime(2026, 6, 9, 21, 30, 15, tzinfo=timezone.utc)
    with patch("src.parsers.admin_inputs._now_utc", return_value=fixed):
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.max_row == 2  # header + one row only


def test_record_change_does_not_dedup_different_cells():
    wb = _new_workbook_with_some_sheets()
    fixed = datetime(2026, 6, 9, 21, 30, 15, tzinfo=timezone.utc)
    with patch("src.parsers.admin_inputs._now_utc", return_value=fixed):
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
        record_change(wb, "Sales Goals", "C5", 50, 75, user="hannah")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.max_row == 3  # header + two rows


def test_record_change_does_not_dedup_same_cell_different_new_value():
    """Same cell edited twice to different values in the same second = 2 rows."""
    wb = _new_workbook_with_some_sheets()
    fixed = datetime(2026, 6, 9, 21, 30, 15, tzinfo=timezone.utc)
    with patch("src.parsers.admin_inputs._now_utc", return_value=fixed):
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
        record_change(wb, "Sales Goals", "B5", 200, 300, user="hannah")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.max_row == 3


def test_record_change_allows_same_row_in_different_seconds():
    """The dedup key is per-second; advance the clock, allow the row again."""
    wb = _new_workbook_with_some_sheets()
    t1 = datetime(2026, 6, 9, 21, 30, 15, tzinfo=timezone.utc)
    t2 = datetime(2026, 6, 9, 21, 30, 16, tzinfo=timezone.utc)
    with patch("src.parsers.admin_inputs._now_utc", side_effect=[t1, t2]):
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.max_row == 3


# ---------------------------------------------------------------------------
# Save / reload round-trip
# ---------------------------------------------------------------------------

def test_change_log_survives_save_reload(tmp_path: Path):
    """End-to-end: write workbook, save, reopen — Change Log row is still there."""
    wb = _new_workbook_with_some_sheets()
    record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah", reason="audit test")
    out = tmp_path / "admin_inputs.xlsx"
    wb.save(out)

    wb2 = openpyxl.load_workbook(out)
    assert CHANGE_LOG_SHEET in wb2.sheetnames
    ws = wb2[CHANGE_LOG_SHEET]
    headers = [ws.cell(row=1, column=c + 1).value for c in range(len(CHANGE_LOG_HEADERS))]
    assert headers == list(CHANGE_LOG_HEADERS)
    assert ws.cell(row=2, column=2).value == "hannah"
    assert ws.cell(row=2, column=3).value == "Sales Goals"
    assert ws.cell(row=2, column=4).value == "B5"
    assert ws.cell(row=2, column=6).value == 200
    assert ws.cell(row=2, column=7).value is None


def test_change_log_appends_across_sessions(tmp_path: Path):
    """Open, record, save, reopen, record again — rows accumulate."""
    out = tmp_path / "admin_inputs.xlsx"

    wb1 = _new_workbook_with_some_sheets()
    record_change(wb1, "Sales Goals", "B5", 100, 200, user="hannah")
    wb1.save(out)

    wb2 = openpyxl.load_workbook(out)
    record_change(wb2, "Sales Goals", "B6", 50, 75, user="jaima")
    wb2.save(out)

    wb3 = openpyxl.load_workbook(out)
    ws = wb3[CHANGE_LOG_SHEET]
    assert ws.max_row == 3  # header + 2 rows


# ---------------------------------------------------------------------------
# Timestamp lives inside the function, not at module load
# ---------------------------------------------------------------------------

def test_timestamp_is_fetched_at_call_time_not_import_time():
    """
    Regression guard: record_change must call datetime.now(timezone.utc)
    INSIDE the function. Otherwise mocking it on a per-test basis (or freezing
    time in snapshot tests) breaks.
    """
    wb = _new_workbook_with_some_sheets()
    fixed = datetime(2026, 6, 9, 21, 30, 15, tzinfo=timezone.utc)
    with patch("src.parsers.admin_inputs._now_utc", return_value=fixed):
        record_change(wb, "Sales Goals", "B5", 100, 200, user="hannah")

    ws = wb[CHANGE_LOG_SHEET]
    assert ws.cell(row=2, column=1).value == fixed.isoformat()
