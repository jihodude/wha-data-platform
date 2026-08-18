"""The receipts window must agree with the program — tested on live-shaped raw.

Fixtures mirror the live shapes from this week's cases (TownePlace late payer,
Governor comp, a clean payer) so the renderer's dispositions are pinned to the
same rulings the pipeline implements.
"""
from reports.membership_performance_tracker.logic.receipts_render import (
    render_retention, gate_retention)

RAW = {"families": {
    "retention": {
        "U1|4": {
            "manager": "U1", "bill_month": "4", "as_of": "2026-05-31",
            "by_account": {
                "ACLEAN": [{"balance": 0.0, "invoice_total": 730.0,
                            "invoice_number": "I1", "invoice_date": "a", "adjustments": 0.0}],
                "ALATE":  [{"balance": 0.0, "invoice_total": 1178.0,
                            "invoice_number": "I2", "invoice_date": "a", "adjustments": 0.0}],
                "ACOMP":  [{"balance": 0.0, "invoice_total": 1891.0, "_comp": True,
                            "invoice_number": "I3", "invoice_date": "a", "adjustments": 0.0},
                           {"balance": 0.0, "invoice_total": -1891.0, "_comp": True,
                            "invoice_number": "I4", "invoice_date": "b", "adjustments": 0.0}],
            },
            "removed_by": {"AGONE": "ask fully rescinded, never reissued"},
        }},
    "retention_payments": {"4": {"I1": "2026-05-10", "I2": "2026-06-10"}},
}}


def _render():
    return render_retention(
        RAW, {"U1": "Snohomish"},
        band_of=lambda t, a: "Target",
        name_of=lambda a: f"name-{a}")


def test_dispositions_match_the_rulings():
    rows = {r["member_id"]: r for r in _render()}
    assert rows["ACLEAN"]["qualifies"] == "paid in window" and rows["ACLEAN"]["counted"]
    assert rows["ACLEAN"]["retained"] == 730.0
    late = rows["ALATE"]
    assert late["qualifies"] == "paid after close" and not late["counted"]
    assert late["retained"] == 0.0, "dollars follow the close (7/31 ruling)"
    assert late["ask"] == 1178.0, "the ask never moves"
    comp = rows["ACOMP"]
    assert comp["qualifies"] == "comped" and comp["counted"]
    assert (comp["ask"], comp["retained"]) == (0.0, 0.0), "$0/$0 (7/31 ruling)"
    assert rows["AGONE"]["qualifies"] == "bill rescinded" and not rows["AGONE"]["counted"]
    assert all(r["month"] == "May" for r in rows.values()), "bm4 files under its close month"


def test_gate_passes_when_receipts_equal_published():
    rows = _render()
    cache = {"ret_paid_target": {"Snohomish": {"May": 2}},      # clean + comp
             "ret_billed_target": {"Snohomish": {"May": 3}},    # + late payer
             "rev_up_target": {"Snohomish": {"May": 1908.0}},   # 730 + 1178
             "rev_retained_target": {"Snohomish": {"May": 730.0}}}
    assert gate_retention(rows, cache) == []


def test_gate_fails_loud_on_any_drift():
    rows = _render()
    cache = {"ret_paid_target": {"Snohomish": {"May": 3}}}
    errs = gate_retention(rows, cache)
    assert errs and "receipts 2" in errs[0]


def test_gates_do_not_skip_cells_over_a_territory_key_spelling():
    """Overnight review (receipts F3): the renderer names territories from
    the SLX descriptor ('SpokaneNE') while the cache uses the display name
    ('Spokane/NE'). The gates looked up the raw key, missed, and skipped —
    80 retention cells went UNGATED. A mismatch must now be caught."""
    from reports.membership_performance_tracker.logic import receipts_render as rr
    rows = [{"metric": "Retention", "month": "May", "territory": "SpokaneNE",
             "band": "Target", "member": "X", "member_id": "A1", "invoice": "",
             "key_date": "", "ask": 510.0, "retained": 510.0,
             "counted": True, "qualifies": "paid in window"}]
    fields = {"ret_paid_target": {"Spokane/NE": {"May": 99}},
              "ret_billed_target": {"Spokane/NE": {"May": 99}},
              "rev_up_target": {"Spokane/NE": {"May": 0.0}},
              "rev_retained_target": {"Spokane/NE": {"May": 0.0}}}
    errors = rr.gate_retention(rows, fields)
    assert errors, "a real disagreement must not hide behind the key spelling"
    assert any("Spokane" in e for e in errors)


def test_reconciliation_rows_are_text_not_excel_formulas():
    """Overnight workbook review: the tie-out line began with '=', so
    openpyxl stored all 471 of them as FORMULAS and Excel renders #N/A —
    the whole reconciliation layer of the transparency tab was error
    cells in the shipped workbook."""
    from openpyxl import Workbook
    from reports.membership_performance_tracker.logic import receipts_sheet as rs
    wb = Workbook()
    rows = [{"metric": "Retention", "month": "May", "territory": "Pierce",
             "band": "Target", "member": "X", "member_id": "A1", "invoice": "",
             "key_date": "", "ask": 510.0, "retained": 510.0,
             "counted": True, "qualifies": "paid in window"}]
    rs.build_sheet(wb, "2026-08", rows)
    ws = wb["Data - MPR Members"]
    formulas = [c.coordinate for row in ws.iter_rows() for c in row
                if c.data_type == "f"]
    assert not formulas, f"no cell may be a formula: {formulas[:5]}"
    assert any("matches the report" in str(c.value)
               for row in ws.iter_rows() for c in row), "tie-out line still present"


def test_new_members_show_the_real_bill_and_the_real_cash():
    """Shannon, 8/12: "what does Collected None mean?" — the trace page showed
    the PAID amount under 'Asked $' and a literal None under 'Collected $'.
    The receipts now carry both truths separately:

        ask      = the net bill (IN + AD)          — what was asked
        retained = amount − bob                     — cash that arrived

    so a partial payer reads asked 1,070 / collected 500, and a BOB comp reads
    asked 505 / collected 0 with Why='BOB comp'. Amount stays untouched — the
    conservation gate sums it against the published face."""
    from reports.membership_performance_tracker.logic import receipts_render as rr
    raw = {"families": {"new_members": {
        # a partial payer: billed 1,070, paid 500 so far
        "A1|Jun": {"territory": "TKP", "month": "Jun", "band": "target",
                   "amount": 500.0, "bob": 0.0, "counted": True,
                   "first_pay": "2026-06-08", "ask": 1070.0},
        # a BOB comp: billed 505, credited 505, no cash
        "A2|Oct": {"territory": "Snohomish", "month": "Oct", "band": "non_target",
                   "amount": 505.0, "bob": 505.0, "counted": True,
                   "first_pay": "2025-10-16", "ask": 505.0},
        # an OLD capture without the ask field — must not crash, ask stays None
        "A3|Jul": {"territory": "Pierce", "month": "Jul", "band": "target",
                   "amount": 730.0, "bob": 0.0, "counted": True,
                   "first_pay": "2026-07-02"},
    }}}
    rows = {r["member_id"]: r for r in rr.render_new_members(raw, lambda a: a)}
    partial = rows["A1"]
    assert partial["ask"] == 1070.0
    assert partial["retained"] == 500.0
    assert partial["amount"] == 500.0          # the gate's field, unchanged
    comp = rows["A2"]
    assert comp["ask"] == 505.0
    assert comp["retained"] == 0.0             # amount − bob: no cash arrived
    assert comp["qualifies"] == "BOB comp"
    old = rows["A3"]
    assert old["ask"] is None                  # old capture: no fabricated ask
    assert old["retained"] == 730.0


def test_drops_receipts_mirror_the_admin_reasons_rule():
    """C1a root cause (cloud log, 8/12): feed_drops excludes admin/record-
    housekeeping reasons (A2, ruled), but render_drops never learned the rule
    — so receipts listed Cheeky Noodles and Gill Brothers as counted drops,
    disagreed with the published face by exactly those rows, and the
    conservation gate refused to publish receipts two runs in a row. The
    renderer must give them an excluded row with the reason, same as the feed."""
    from types import SimpleNamespace
    from reports.membership_performance_tracker.logic import receipts_render as rr

    def _row(name, reason, status="Inactive", bm=4):
        return SimpleNamespace(name=name, mid="0", territory="Southwest",
                               dues=730.0, bill_month=bm, status=status,
                               status_reason=reason, business_type="",
                               status_date=None)

    admin = _row("Cheeky Noodles", "New Lead")
    real = _row("Real Loss LLC", "Non-Payment")
    out = rr._classify_drop_rows_for_test([admin, real], period="2026-08")
    by = {r["member"]: r for r in out}
    assert by["Cheeky Noodles"]["counted"] is False
    why = by["Cheeky Noodles"]["qualifies"]
    # The CRM's own reason is kept verbatim (closed vocabulary); what we add is
    # the plain-English explanation of why it doesn't count. Assert both, not
    # the exact prose — this line was pinned to "housekeeping" and broke when
    # the wording was made readable (8/13).
    assert why.startswith("New Lead"), "the CRM's own reason must survive"
    assert "not counted as a drop" in why.lower(), (
        "an excluded row must say, in plain words, that it isn't a member loss")
    assert by["Real Loss LLC"]["counted"] is True


# ── THE T/NT LOCK ────────────────────────────────────────────────────────────
# Allied is Non-Target by definition (ruled 2026-08-11) and the face applies
# that lock in `_retention_target_split`. These receipts have to agree, or the
# conservation gate refuses the cell.
#
# Found 2026-08-13, in July: the renderer took its band from the target map,
# which had banded Seattle Laundry on size, so the receipt said "Target" while
# the amendment moved a NON-target cell. band_of returns "Target" here for
# exactly that reason — the lock must beat it.

ALLIED_RAW = {"families": {
    "retention": {
        "U1|6": {
            "manager": "U1", "bill_month": "6", "as_of": "2026-07-31",
            "by_account": {
                "AALLIED": [{"balance": 0.0, "invoice_total": 505.0, "_allied": True,
                             "invoice_number": "I9", "invoice_date": "a",
                             "adjustments": 0.0, "dues_level": 505.0}],
                "AHOSP":   [{"balance": 0.0, "invoice_total": 730.0,
                             "invoice_number": "I8", "invoice_date": "a",
                             "adjustments": 0.0}],
            },
            "removed_by": {},
        }},
    "retention_payments": {"6": {"I9": "2026-07-10", "I8": "2026-07-10"}},
}}


def test_allied_is_non_target_however_the_target_map_bands_it():
    rows = {r["member_id"]: r for r in render_retention(
        ALLIED_RAW, {"U1": "Snohomish"},
        band_of=lambda t, a: "Target",          # the map would promote it
        name_of=lambda a: f"name-{a}")}
    assert rows["AALLIED"]["band"] == "Non-Target", (
        "allied is Non-Target by definition — size must never promote it, or "
        "the receipts disagree with the face and conservation refuses the cell")
    assert rows["AHOSP"]["band"] == "Target", "the lock must not touch hospitality"
    assert rows["AALLIED"]["counted"] and rows["AALLIED"]["retained"] == 505.0, (
        "the lock changes the band only — never whether it counted or for how much")


# ── C10: the published filing wins over a render-time re-derivation ──────────
# Two classes never file at their own date — dateless drops go to the class
# close, billing-stale ones to the status-flip date — and feed_drops records
# each decision as it makes it. build_receipts never handed those decisions to
# the renderer, so it re-derived them and landed a month later than the face
# (measured 8/12: receipts Dec = face Nov + face Dec across six territories).
# Those cells then read as conservation mismatches on every build.

def test_filed_overrides_map_is_built_from_the_feeds_own_decisions():
    from scripts.build_receipts import filed_overrides_from_flags
    flags = {
        "drop_date_fallback": [
            {"name": "K.D. Grazie, Inc.", "mid": "0053128", "bill_month": 6,
             "filed": "Jul", "reason": "no end date on record at capture"}],
        "hygiene_counted": [
            {"name": "Fogo De Chao", "mid": "0066309", "filed": "Feb",
             "defect": "counted at status-flip date; billing is stale"}],
        "no_cycle": [{"name": "Brinker International", "mid": "0025225"}],
    }
    ov = filed_overrides_from_flags(flags)
    assert ov["0053128"][0] == "Jul", "dateless drops file at the class close"
    assert ov["0066309"][0] == "Feb", "billing-stale drops file at the flip date"
    assert "0025225" not in ov, "a flag with no filing decision must not override"
    assert all(note for _, note in ov.values()), (
        "every override carries a note — a moved filing must never be silent")


def test_a_published_filing_overrides_the_renderers_own_month():
    rows = render_retention(RAW, {"U1": "Snohomish"},
                            band_of=lambda t, a: "Target",
                            name_of=lambda a: a)
    assert rows, "sanity: the retention fixture still renders"
    # render_drops applies overrides by mid; prove the contract it relies on
    from scripts.build_receipts import filed_overrides_from_flags
    ov = filed_overrides_from_flags(
        {"drop_date_fallback": [{"mid": "0001", "filed": "Jul", "reason": "x"}]})
    assert ov.get("0001") == ("Jul", "date estimated — filed as published")


# ── Nobody falls off the table ───────────────────────────────────────────────
# A fresh row can file INTO a closed month — a dateless drop lands at its class
# close, or a CRM edit (or one of our own fixes) moves it. The closed edition
# was photographed earlier and has never heard of it, so it used to be dropped
# on both sides at once: gone from the open month it left, absent from the
# closed month it arrived in, counted nowhere. Bob's Burgers & Brew corporate
# ($2,775) and Marina Square ($510) went missing exactly this way on 8/13, and
# nothing said so — the cell gate compares sums, and zero matched zero.

def _sealed_inputs():
    carried = [{"metric": "Drops", "month": "Jul", "territory": "TKP",
                "member": "Hungry Harbor", "member_id": "0011", "counted": True,
                "band": "Target", "amount": 700.0, "qualifies": "Non-Payment"}]
    return ["Jul"], {("Drops", "Jul"): carried}


def test_a_row_landing_in_a_closed_month_is_kept_and_explained():
    from scripts.build_receipts import merge_sealed
    sealed_months, carried = _sealed_inputs()
    fresh = [
        {"metric": "Drops", "month": "Jul", "territory": "NorthCentral",
         "member": "Bob's Burgers & Brew", "member_id": "0056458", "counted": True,
         "band": "Target", "amount": 2775.0, "qualifies": "Non-Payment"},
        {"metric": "Drops", "month": "Feb", "territory": "NorthCentral",
         "member": "Tino's", "member_id": "0022", "counted": True,
         "band": "Target", "amount": 500.0, "qualifies": "Sold"},
    ]
    out = merge_sealed(fresh, "Drops", sealed_months, carried, fresh_families=set())
    bobs = [r for r in out if r["member_id"] == "0056458"]
    assert bobs, "a row filing into a closed month must not vanish"
    assert bobs[0]["counted"] is False, (
        "it is not in the closed month's locked numbers, so it must not count")
    assert "closed" in bobs[0]["qualifies"], "it must say why it isn't counted"
    assert bobs[0]["month"] == "Jul", "it belongs under the month it arrived in"
    assert any(r["member_id"] == "0011" for r in out), "carried rows survive"
    assert any(r["member_id"] == "0022" for r in out), "open months render fresh"


def test_a_member_already_in_the_closed_edition_is_not_duplicated():
    from scripts.build_receipts import merge_sealed
    sealed_months, carried = _sealed_inputs()
    fresh = [{"metric": "Drops", "month": "Jul", "territory": "TKP",
              "member": "Hungry Harbor", "member_id": "0011", "counted": True,
              "band": "Target", "amount": 700.0, "qualifies": "Non-Payment"}]
    out = merge_sealed(fresh, "Drops", sealed_months, carried, fresh_families=set())
    assert len([r for r in out if r["member_id"] == "0011"]) == 1, (
        "the carried row is the record — a fresh copy must not double it")


def test_counted_governs_the_count_and_money_governs_the_dollars():
    """Settled 2026-08-20, after each half of the old conflation bit in turn.

    2026-08-14: a $575 sale stranded into sealed July made the gate refuse
    every receipt — fixed then by skipping uncounted rows entirely.
    2026-08-18/19/20: a $1,754.50 REFUND — an uncounted row whose money is
    real and negative — sat inside the published NorthKing Aug cell while the
    gate skipped it, blocking receipts three nights running over a number
    that was CORRECT.

    The two flags mean different things. `counted` = in the member-count
    metric. Money = the row's amount, summed over EVERY row — and a row that
    must move no total (the stranded case) carries amount 0 explicitly
    (merge_sealed zeroes it, the same convention as retention's exclusions).
    """
    from reports.membership_performance_tracker.logic import receipts_render as rr

    def row(amount, counted, bob=0.0):
        return {"metric": "New Members / New Sales", "territory": "NorthKing",
                "month": "Aug", "band": "Target", "amount": amount,
                "bob": bob, "counted": counted, "qualifies": "first payment"}

    published = {"revenue_target": {"NorthKing": {"Aug": 2924.0}},
                 "new_members_target": {"NorthKing": {"Aug": 2}},
                 "revenue_bob_target": {"NorthKing": {"Aug": 0.0}}}

    # The live 8/18 shape: two counted sales + one uncounted refund.
    rows = [row(3795.0, True), row(883.5, True), row(-1754.5, False)]
    assert rr.gate_new_members(rows, published) == [], (
        "an uncounted refund's money is real — the dollars must include it, "
        "and the count must not")

    # A stranded row carries amount 0 (merge_sealed's job) — moves nothing.
    rows_stranded = rows + [row(0.0, False)]
    assert rr.gate_new_members(rows_stranded, published) == []

    # And the gate still bites: a counted row that disagrees must fail.
    assert rr.gate_new_members([row(3795.0, True)], published), \
        "this fix must not turn the gate off"


def test_merge_sealed_strands_carry_no_money():
    """The stranded row's amount is zeroed ON THE ROW (with the figure kept
    visible in its text), because the dollars gates sum every row."""
    import importlib
    br = importlib.import_module("scripts.build_receipts")
    fresh = [{"metric": "New Members / New Sales", "month": "Jul",
              "territory": "Spokane/NE", "band": "Non-Target",
              "member": "M", "member_id": "A1", "amount": 575.0, "bob": 5.0,
              "counted": True, "qualifies": "first payment"}]
    out = br.merge_sealed(fresh, "New Members / New Sales", {"Jul"}, {}, set())
    stranded = [r for r in out if not r["counted"]]
    assert len(stranded) == 1
    s = stranded[0]
    assert s["amount"] == 0.0 and s["bob"] == 0.0
    assert "575" in s["qualifies"] and "closed" in s["qualifies"]
