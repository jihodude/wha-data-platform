"""Defect #5 (meaning audit 2026-07-17): the "Billables - NRA (#)" report row is
removed. The bm>12 NRA marker was DISPROVEN 2026-07-14 (Duesbillmonth is stamped
on every location record); NRA chain accounts are counted inside the Majors/NRA
territory column, so a dedicated NRA row was a permanent, misleading 0.
Warrant: DECISIONS.md 2026-07-14 (bill-month rule reverted) + REFERENCE.md
(corrected 2026-07-17).
"""
from pathlib import Path

import openpyxl

from reports.membership_performance_tracker.mapper import METRIC_FIELD_MAP

TEMPLATE = Path("reports/membership_performance_tracker/template.xlsx")


def test_nra_metric_absent_from_mapper():
    assert "Billables - NRA (#)" not in METRIC_FIELD_MAP


def test_nra_metric_absent_from_ref_metrics_catalog():
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True)
    ws = wb["Ref - Metrics"]
    names = {ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)}
    assert "Billables - NRA (#)" not in names
    # the other billables metrics remain
    assert "Billables - Allied (#)" in names
