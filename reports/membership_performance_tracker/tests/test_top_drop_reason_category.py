"""Top Drop Reason — RETIRED (8/6/2026 ruling, Jiho): the column could only
ever say Non-Payment or Sold (the pipeline emits per-REASON lines, never the
4-category rollups the old formula spoke, so 'Out of Business' — the real #1
— was structurally invisible) and a single mode "doesn't give any insights…
they can just go look" at the Data - Drops detail. This test now pins the
retirement: the column must stay gone from the dashboard drops section.

History: the original test (meaning-audit defect #3, 2026-07-17) pinned the
formula to category metrics; superseded by the retirement ruling.
"""
from pathlib import Path

import openpyxl

TEMPLATE = Path("reports/membership_performance_tracker/template.xlsx")


def test_top_drop_reason_column_stays_retired():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Summary Dashboard"]
    labels = {str(ws.cell(r, c).value or "")
              for r in range(1, ws.max_row + 1) for c in range(1, 12)}
    assert "Top Drop Reason" not in labels
    header = next(r for r in range(1, ws.max_row + 1)
                  if "DROPS — Current Month Detail"
                  in str(ws.cell(r, 1).value or ""))
    for r in range(header + 1, header + 13):
        assert ws.cell(r, 8).value is None, f"H{r} should be empty (retired)"
