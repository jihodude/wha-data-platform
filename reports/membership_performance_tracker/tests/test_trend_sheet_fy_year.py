"""Defect #4 (meaning audit): Monthly Trends 'Oct/Nov/Dec' columns must query the
FISCAL-year-start calendar year (=$B$4-1), matching the TM sheets and how
Data-Monthly-Metrics stores those rows (Oct-Dec 2025 for FY2025-26). The trend
sheet was querying $B$4 for every month, so Oct-Dec resolved to a future year and
returned 0, and Season Total silently omitted the fiscal Q1.
Warrant: AGENTS.md hard rule (FY = Oct 1-Sep 30) + TM sheets already use $B$4-1.
"""
from reports.membership_performance_tracker.trend_sheet import statewide_sumifs


def test_oct_nov_dec_query_prior_fiscal_year():
    for m in ("Oct", "Nov", "Dec"):
        f = statewide_sumifs(m, "New Members (#)")
        assert "$B$4-1" in f, f"{m} must query $B$4-1 (fiscal-year start); got: {f}"


def test_jan_to_sep_query_selector_year():
    for m in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"):
        f = statewide_sumifs(m, "New Members (#)")
        assert ",$B$4," in f, f"{m} must query $B$4; got: {f}"
        assert "$B$4-1" not in f, f"{m} must NOT subtract a year; got: {f}"


# ---------------------------------------------------------------------------
# S3 hole 3 (overnight review 8/11): the two sheets mean different things by
# "B4" and the trends page mirrors the dashboard's value without converting.
#
#   Summary Dashboard B4 = the CALENDAR year of the month in E4 (correct for
#                          it — B4 + E4 together name one real month).
#   Monthly Trends    B4 = the FISCAL-year END year — which is precisely why
#                          Oct/Nov/Dec query $B$4-1 above.
#
# Equal for Jan-Sep, off by one for Oct-Dec. Every period ever run (2026-03 …
# 2026-08) is Jan-Sep, so it has never fired. It fires on the FIRST October
# build and renders the ENTIRE page as the PRIOR fiscal year — and the data
# sheet keeps history (2025: 1,476 rows, 2026: 4,753 in the 8/10 artifact), so
# it shows real, fully-populated last-year numbers rather than an obvious zero.
# Wrong for the Oct/Nov/Dec editions, then silently correct again in January.
# ---------------------------------------------------------------------------
import openpyxl
from reports.membership_performance_tracker.trend_sheet import (
    _default_year, SUMMARY_SHEET)


def _dash(b4, e4):
    """A workbook carrying just the dashboard's two selector cells."""
    wb = openpyxl.Workbook()
    wb.active.title = SUMMARY_SHEET
    wb[SUMMARY_SHEET]["B4"] = b4
    wb[SUMMARY_SHEET]["E4"] = e4
    return wb


def test_october_dashboard_rolls_the_trends_page_to_the_next_fiscal_year():
    """Oct 2026 is FY2026-27. The dashboard says 2026 (October's calendar
    year); the trends selector must be 2027 so Oct-Dec resolve to 2026."""
    assert _default_year(_dash(2026, "Oct")) == 2027


def test_nov_and_dec_dashboards_roll_the_same_way():
    assert _default_year(_dash(2026, "Nov")) == 2027
    assert _default_year(_dash(2026, "Dec")) == 2027


def test_jan_to_sep_dashboard_is_already_the_fiscal_year_end():
    """Mar 2026 is FY2025-26: the calendar year IS the fiscal-year end, so the
    mirrored value must pass through untouched (today's shipped behavior)."""
    for m in ("Jan", "Mar", "Jun", "Aug", "Sep"):
        assert _default_year(_dash(2026, m)) == 2026, m


def test_the_period_fallback_converts_october_too(monkeypatch):
    """Same defect in the no-dashboard path: it takes MPR_PERIOD's calendar
    year and ignores the month entirely."""
    monkeypatch.setenv("MPR_PERIOD", "2026-10")
    assert _default_year(openpyxl.Workbook()) == 2027


def test_the_period_fallback_leaves_a_spring_period_alone(monkeypatch):
    monkeypatch.setenv("MPR_PERIOD", "2026-03")
    assert _default_year(openpyxl.Workbook()) == 2026
