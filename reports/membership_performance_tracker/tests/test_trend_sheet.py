"""
test_trend_sheet.py — All-TM monthly trend page (exec feedback #3 / master-list #5).

A new sheet inserted between "Summary Dashboard" and the first TM sheet, showing
every metric SUMMED across all territories (statewide) for every month, so execs
can read month-over-month trends in number format.

Design (mirrors the formula-driven Summary Dashboard):
  - The page is pure Excel formulas reading from "Data - Monthly Metrics".
  - Additive metrics (count/currency) → statewide SUMIFS per month (NO territory
    criterion → sums all territories).
  - Percentage metrics are NOT naively summed (you can't add percentages). They
    are computed as statewide ratios from the additive component rows already on
    the page (e.g. Penetration Combined = ΣTargetMembers / ΣTargetLocations).

These tests assert the deterministic formula strings + sheet placement, since no
headless Excel is available to evaluate the workbook.
"""

from pathlib import Path

import openpyxl
import pytest

from reports.membership_performance_tracker.trend_sheet import (
    build_monthly_trend_sheet,
    TREND_SHEET_NAME,
    DATA_SHEET,
    MONTHS_FISCAL,
)

TEMPLATE = Path("reports/membership_performance_tracker/template.xlsx")


@pytest.fixture
def wb():
    # Real template = highest-fidelity fixture (real metric catalog, real sheet order).
    # Loaded fresh per test; never saved back over the template.
    return openpyxl.load_workbook(TEMPLATE)


def test_inserts_after_summary_dashboard(wb):
    meta = build_monthly_trend_sheet(wb)
    names = wb.sheetnames
    summary_idx = names.index("Summary Dashboard")
    assert names[summary_idx + 1] == TREND_SHEET_NAME
    # the sheet that used to follow Summary Dashboard is now pushed one slot down
    assert meta["sheet_name"] == TREND_SHEET_NAME


def test_header_months_are_fiscal_order(wb):
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    hr = meta["header_row"]
    # Column B..M hold the 12 fiscal months Oct..Sep
    assert MONTHS_FISCAL == ["Oct", "Nov", "Dec", "Jan", "Feb", "Mar",
                             "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
    months_in_header = [ws.cell(row=hr, column=2 + i).value for i in range(12)]
    assert months_in_header == MONTHS_FISCAL
    # Metric label column header + a season total column after the months
    assert ws.cell(row=hr, column=1).value == "Metric"
    assert ws.cell(row=hr, column=14).value == "Season Total"


def test_additive_metric_uses_statewide_sumifs(wb):
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    row = meta["metric_rows"]["Billables - Allied (#)"]
    # Oct is the first month column (B). Oct/Nov/Dec query the fiscal-year START
    # calendar year ($B$4-1) — Data-Monthly-Metrics stores them under that year
    # (meaning-audit fix 2026-07-17; was wrongly $B$4, so Oct-Dec resolved to 0).
    formula = ws.cell(row=row, column=2).value
    expected = (
        f"=SUMIFS('{DATA_SHEET}'!E:E,"
        f"'{DATA_SHEET}'!A:A,$B$4-1,"
        f"'{DATA_SHEET}'!B:B,\"Oct\","
        f"'{DATA_SHEET}'!D:D,\"Billables - Allied (#)\")"
    )
    assert formula == expected
    # No territory criterion anywhere → it sums across every TM
    assert "C:C" not in formula


def test_season_cell_sums_flows_and_shows_the_latest_photo_for_stocks(wb):
    """A STOCK IS NEVER SUMMED (ruled 8/6). This test previously asserted
    the opposite for Billables and so enshrined the defect the overnight
    review found: with prior photo months populated, the season cell read
    roughly twice the roster and grew every month."""
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    stock = meta["metric_rows"]["Billables - Allied (#)"]
    assert ws.cell(row=stock, column=14).value == (
        f"=IFERROR(LOOKUP(2,1/(B{stock}:M{stock}<>0),B{stock}:M{stock}),0)")
    flow = meta["metric_rows"]["New Sales Revenue - Target ($)"]
    assert ws.cell(row=flow, column=14).value == f"=SUM(B{flow}:M{flow})"


def test_percentage_metrics_not_present_as_additive_rows(wb):
    meta = build_monthly_trend_sheet(wb)
    # The naive pct metrics must NOT appear as summed rows.
    for pct_metric in [
        "Penetration - Combined %",
        "Penetration - Restaurant %",
        "Penetration - Lodging %",
        "Goal - Retention %",
    ]:
        assert pct_metric not in meta["metric_rows"]


def test_combined_penetration_is_computed_ratio(wb):
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    rate_row = meta["rate_rows"]["Penetration - Combined %"]
    mr = meta["metric_rows"]
    tm_r = mr["Target Members - Restaurant (#)"]
    tm_l = mr["Target Members - Lodging (#)"]
    tl_r = mr["Total Target Locations - Restaurant (#)"]
    tl_l = mr["Total Target Locations - Lodging (#)"]
    # Oct column (B): ratio of summed members to summed locations, guarded /0
    formula = ws.cell(row=rate_row, column=2).value
    expected = (
        f"=IF((B{tl_r}+B{tl_l})=0,0,"
        f"(B{tm_r}+B{tm_l})/(B{tl_r}+B{tl_l}))"
    )
    assert formula == expected


def test_idempotent_rebuild_does_not_duplicate(wb):
    build_monthly_trend_sheet(wb)
    n1 = wb.sheetnames.count(TREND_SHEET_NAME)
    build_monthly_trend_sheet(wb)
    n2 = wb.sheetnames.count(TREND_SHEET_NAME)
    assert n1 == 1
    assert n2 == 1


def test_retention_and_penetration_rates_carry_goal_colors(wb):
    """8/10 (Jiho): color the retention + penetration rate rows against the
    admin-input goals, dynamically — same green/yellow/red language as the
    TM sheets. Goals come from the data sheet (AVERAGEIFS across
    territories), guarded so no-goal months stay uncolored."""
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    want_rows = {meta["rate_rows"]["Retention % — T + NT (#)"],
                 meta["rate_rows"]["Penetration - Combined %"]}
    texts = []
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for cr in rng.sqref.ranges:
            if cr.min_row in want_rows and cr.max_row == cr.min_row:
                for rule in rules:
                    texts.append(" ".join(rule.formula or []))
    joined = " ".join(texts)
    assert "AVERAGEIFS" in joined, "goal must come from the data dynamically"
    assert "Goal - Retention %" in joined
    assert "Goal - Penetration %" in joined
    assert "COUNTIFS" in joined, "months with no goal entered stay uncolored"
    assert ")>0" in joined, "months with an empty denominator stay uncolored"


def test_goal_color_rules_are_single_cell_for_excel_online(wb):
    """8/12 (Jiho's screenshot): SharePoint's Excel Online evaluates only the
    ANCHOR cell of a ranged expression rule — Oct and Jan colored, Dec at
    97.2% vs a 92% goal stayed bare. Every goal rule must therefore target
    exactly one cell, with that cell's own references baked in."""
    meta = build_monthly_trend_sheet(wb)
    ws = wb[TREND_SHEET_NAME]
    checked = 0
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        joined = " ".join(f for rule in rules for f in (rule.formula or []))
        if "Goal - Retention %" not in joined and \
                "Goal - Penetration %" not in joined:
            continue
        for cr in rng.sqref.ranges:
            assert (cr.min_row == cr.max_row and cr.min_col == cr.max_col), \
                f"ranged goal rule survives at {cr.coord}"
            # the formula must reference ITS OWN column, not the anchor's
            from openpyxl.utils import get_column_letter
            col = get_column_letter(cr.min_col)
            assert f"{col}{cr.min_row}" in joined, \
                f"rule at {cr.coord} references a different cell: {joined[:90]}"
            checked += 1
    assert checked >= 24, "goal rules must exist for the rate rows"


def test_template_static_sheets_carry_no_ranged_goal_rules():
    """scripts/fix_pct_cf_per_cell.py applied the same per-cell surgery to
    the committed template (TM sheets etc.) — it must never regress."""
    twb = openpyxl.load_workbook(TEMPLATE)
    for ws in twb.worksheets:
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            joined = " ".join(f for rule in rules for f in (rule.formula or []))
            if "Goal - Retention %" not in joined and \
                    "Goal - Penetration %" not in joined:
                continue
            for cr in rng.sqref.ranges:
                assert (cr.min_row == cr.max_row
                        and cr.min_col == cr.max_col), \
                    f"{ws.title}: ranged goal rule at {cr.coord}"


def test_bob_summary_rows_appear_on_the_trends_sheet(wb):
    """Anthony (demo 8/12 22:53): "if it's possible to summarize like all of
    the bobs... there is no bob line here." The catalog now registers the
    four BOB metrics, so the statewide Trends sheet builds their rows."""
    meta = build_monthly_trend_sheet(wb)
    for label in ("New Members - BOB (#)",
                  "New Sales Revenue - BOB/Promo ($)",
                  "Retention - BOB (#)",
                  "Retention - BOB ($)"):
        assert label in meta["metric_rows"], f"missing BOB row: {label}"


def test_template_goal_rules_all_stop_when_they_fire():
    """8/12 night (fresh build, Excel Online): 100% showed YELLOW because the
    TM sheets' goal-rule trios carried stopIfTrue=None — the red catch-all
    fires on every numeric cell, so the trio is only correct under stop
    semantics. Every goal rule must stop."""
    twb = openpyxl.load_workbook(TEMPLATE)
    checked = 0
    for ws in twb.worksheets:
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                joined = " ".join(rule.formula or [])
                if "Goal - Retention %" in joined or \
                        "Goal - Penetration %" in joined:
                    assert rule.stopIfTrue, \
                        f"{ws.title} {rng.sqref}: goal rule without stopIfTrue"
                    checked += 1
    assert checked > 1000


def test_no_duplicate_conditional_format_priorities_anywhere():
    """Shannon, 8/13: Excel repaired the workbook and stripped cell
    information from every TM sheet. A worksheet's conditional-format rules
    must each carry a UNIQUE priority; the 8/13 template surgery rebuilt rule
    lists with deepcopy, which carries the original priority along, leaving
    168 collisions per TM sheet. Excel treats that as corruption."""
    twb = openpyxl.load_workbook(TEMPLATE)
    for ws in twb.worksheets:
        seen = {}
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                p = rule.priority
                assert p not in seen, (
                    f"{ws.title}: priority {p} used twice "
                    f"({seen[p]} and {rng.sqref}) — Excel will repair the file")
                seen[p] = str(rng.sqref)


def test_generated_workbook_keeps_unique_priorities(tmp_path):
    """The guard has to survive the build, not just the template."""
    import re, zipfile
    from collections import Counter
    from reports.membership_performance_tracker.writer import TrackerV4Writer
    out = tmp_path / "w.xlsx"
    TrackerV4Writer(template_path=TEMPLATE).write(out, metric_rows=[])
    z = zipfile.ZipFile(out)
    for name in [n for n in z.namelist()
                 if re.match(r"xl/worksheets/sheet\d+\.xml$", n)]:
        xml = z.read(name).decode("utf8", "replace")
        prios = [int(p) for p in re.findall(r'<cfRule[^>]*priority="(\d+)"', xml)]
        dupes = [p for p, c in Counter(prios).items() if c > 1]
        assert not dupes, f"{name}: duplicate CF priorities {dupes[:5]}"
