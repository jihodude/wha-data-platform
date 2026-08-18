"""The Data - MPR Flags sheet — human layout (Jiho 8/10: colors, consistent
columns, correct info; no raw section keys, no account ids as names)."""
from openpyxl import Workbook

from reports.membership_performance_tracker.logic import receipts_sheet as rs


def _flat(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_flags_sheet_consistent_human_layout():
    wb = Workbook()
    flags = {
        "active_without_payment": [
            {"account_id": "A1", "name": "A1", "territory": "Pierce",
             "month": "Jan", "why": "enrolled and active with no dues "
             "payment — policy says this should never happen"}],
        "duplicate_records": [
            {"name": "Chain Co", "accounts": ["A1", "A2"],
             "defect": "multiple account records share this name"}],
        "uncovered_billed_accounts": [],
    }
    rs.build_flags_sheet(wb, "2026-08", flags, names={"A1": "Real Name LLC"})
    ws = wb["Data - MPR Flags"]
    rows = _flat(ws)
    flat = "\n".join(str(r) for r in rows)
    assert "Real Name LLC" in flat, "account id must resolve to the name"
    assert "active_without_payment" not in flat, "no raw keys in the sheet"
    assert "uncovered" not in flat.lower(), "empty sections are skipped"
    assert any(r[:4] == ["Member", "WRA #", "Territory", "Month"]
               for r in rows), "consistent column header"
    fills = {str(c.fill.start_color.rgb)
             for row in ws.iter_rows() for c in row
             if c.fill is not None and c.fill.fill_type == "solid"}
    assert any(f.endswith("C00000") for f in fills), "severity coloration"


def test_drops_schedule_window_rows_render_as_schedule_level_notes():
    """S3 hole 1 (ruled 8/11): the window problem is not member-shaped — no
    member, no WRA #, no amount. Plain-string rows land in 'What happened' with
    the member columns blank, and the section carries the annual Oct-1 remedy
    so the operator can act without reading any doc.
    """
    from reports.membership_performance_tracker.logic.crystal import (
        DROPS_WINDOW_FLAG)

    wb = Workbook()
    problem = ("Dropped Allied Report: window starts 2025-10-01, "
               "expected 2026-10-01")
    rs.build_flags_sheet(wb, "2026-10", {DROPS_WINDOW_FLAG: [problem]})
    ws = wb["Data - MPR Flags"]
    rows = _flat(ws)
    flat = "\n".join(str(r) for r in rows)

    assert DROPS_WINDOW_FLAG not in flat, "no raw section key on the sheet"
    assert "not scoped to this fiscal year" in flat, "section title renders"
    assert "StatusDate range" in flat, "the remedy is on the sheet"

    hit = [r for r in rows if r and r[-1] and problem in str(r[-1])]
    assert hit, "the problem text lands in the 'What happened' column"
    assert not any(hit[0][:5]), "member columns stay blank for a schedule note"
