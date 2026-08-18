"""The data-defect flag net (scope locked with Jiho 2026-08-04).

Rule-based exclusions (first-cycle, retro, comps) are receipts rows; but a
member the pipeline drops or judges conservatively because the DATA is bad
(missing status, same-day void, short-gap comp credit, non-standard bill
comment that fails the roster test) must emit a FLAG — a silent drop on bad
data is structurally forbidden. Flags land in flags_<period>.json (beside
the receipts JSON on SP) and on a 'Data - MPR Flags' workbook tab.
"""
from datetime import date

import reports.membership_performance_tracker.logic.retention as R
from reports.membership_performance_tracker.tests.test_comp_lines import (
    _CompSLX, _VoidSLX)


def test_voided_bill_emits_a_flag():
    flags = {}
    R.compute_retention_all_months(
        _VoidSLX(), {"U1": "TKP"}, bill_months=[5], fiscal_year_start=2025,
        request_delay=0.0, today=date(2026, 7, 20), flags_out=flags)
    rows = flags.get("voided_bills") or []
    assert len(rows) == 1, f"the same-day reversal must be flagged: {flags}"
    assert rows[0]["account"] == "COMP1"
    assert rows[0]["bill_month"] == 5


def test_status_unknown_conservative_exclusion_emits_a_flag():
    class _NoStatusSLX(_CompSLX):
        def _fetch_all(self, path, where="", **kw):
            if path == "accounts" and "CDiverseOwnership" not in where:
                return []          # SLX has no statuses — incomplete data
            return super()._fetch_all(path, where=where, **kw)

    flags = {}
    R.compute_retention_all_months(
        _NoStatusSLX(), {"U1": "TKP"}, bill_months=[5], fiscal_year_start=2025,
        request_delay=0.0, today=date(2026, 7, 20), flags_out=flags)
    rows = flags.get("status_unknown_excluded") or []
    assert any(r["account"] == "COMP1" for r in rows), \
        f"conservative exclusion on missing status must be visible: {flags}"


def test_short_gap_comp_emits_review_flag():
    class _ShortGapSLX(_CompSLX):
        def _fetch_all(self, path, where="", **kw):
            rows = super()._fetch_all(path, where=where, **kw)
            if path == "dlInvoiceHistoryHeader":
                for r in rows:
                    if r["Invoice_number"] == "I-COMPC":
                        # credit 3 days after the bill — comp, but review
                        r["Invoice_Date"] = "/Date(1745259200000)/"
            return rows

    flags = {}
    res = R.compute_retention_all_months(
        _ShortGapSLX(), {"U1": "TKP"}, bill_months=[5], fiscal_year_start=2025,
        request_delay=0.0, today=date(2026, 7, 20), flags_out=flags)
    assert res[5]["TKP"]["comped_count"] == 1, "short gap is still a comp"
    rows = flags.get("comp_review") or []
    assert any(r["account"] == "COMP1" for r in rows), \
        f"1-6 day credit gap must be flagged for review: {flags}"


def test_flags_sheet_renders_every_family():
    """'Data - MPR Flags' tab: one section per flag family, one row per
    flagged record. 8/10 redesign: human section titles (no raw keys),
    consistent columns; unknown families still render all their data."""
    from openpyxl import Workbook
    from reports.membership_performance_tracker.logic import receipts_sheet as rs

    flags = {
        "voided_bills": [{"account": "A1", "bill_month": 5,
                          "defect": "same-day reversal"}],
        "comp_review": [{"account": "A2", "bill_month": 6,
                         "defect": "3-day credit gap"}],
        "billable_without_invoice": [{"member": "Four Points", "wra": "123"}],
        "_carried_from_previous_run": ["active_without_payment"],   # meta: skipped
    }
    wb = Workbook()
    rs.build_flags_sheet(wb, "2026-07", flags)
    assert "Data - MPR Flags" in wb.sheetnames
    ws = wb["Data - MPR Flags"]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Voided bills" in text and "Comp review" in text
    assert "voided_bills" not in text, "raw keys stay internal (8/10)"
    assert "same-day reversal" in text and "Four Points" in text
    assert "_carried_from_previous_run" not in text, "meta keys stay internal"


def test_flags_sheet_replaces_itself_on_rebuild():
    from openpyxl import Workbook
    from reports.membership_performance_tracker.logic import receipts_sheet as rs
    wb = Workbook()
    rs.build_flags_sheet(wb, "2026-07", {"a": [{"x": 1}]})
    rs.build_flags_sheet(wb, "2026-07", {"b": [{"y": 2}]})
    assert [s for s in wb.sheetnames].count("Data - MPR Flags") == 1


def test_coverage_leftovers_classifies_the_uncovered():
    """CONSERVATION (scope locked 8/4): every account billed this FY must
    exit the pipeline with a receipts disposition or a flag. Anyone left is
    a silent exclusion by definition — classified so only true unknowns
    scream."""
    from reports.membership_performance_tracker.logic.receipts_render import (
        coverage_leftovers)

    invoiced = {"COVERED", "FLAGGED", "RETRO1", "NOMGR", "MYSTERY"}
    receipts_rows = [{"member_id": "COVERED", "metric": "Retention",
                      "counted": True}]
    flags = {"voided_bills": [{"account": "FLAGGED", "bill_month": 5}]}
    info = {
        "RETRO1":  {"status": "RRO", "manager": "U1"},
        "NOMGR":   {"status": "Active", "manager": None},
        "MYSTERY": {"status": "Active", "manager": "U1"},
    }
    out = coverage_leftovers(invoiced, receipts_rows, flags,
                             account_info=info)
    by_acct = {r["account"]: r for r in out}
    assert "COVERED" not in by_acct and "FLAGGED" not in by_acct
    assert by_acct["RETRO1"]["class"] == "retro/NRA book (excluded by rule)"
    assert by_acct["NOMGR"]["class"] == "no account manager (unassigned book)"
    assert by_acct["MYSTERY"]["class"] == "UNKNOWN — investigate"


def test_duplicate_names_flags_multi_record_families():
    """Jiho's ruling (8/4 evening): duplicates are FLAGGED, humans decide.
    Same normalized name on multiple account ids touching FY money = the
    class that broke name-matching (Shari's, Four Points, Sun Mountain)."""
    from reports.membership_performance_tracker.logic.receipts_render import (
        duplicate_names)
    names = {"A1": "Four Points by Sheraton Bellingham",
             "A2": "Four Points By Sheraton Bellingham ",
             "A3": "Lucky Eagle Casino Hotel"}
    out = duplicate_names(names)
    assert len(out) == 1
    assert set(out[0]["accounts"]) == {"A1", "A2"}
    assert "Four Points" in out[0]["name"]


def test_coverage_output_families_never_self_cover():
    """8/10: the conservation check's own output families must not count as
    coverage on the next run — they made the uncovered count alternate
    between 0 and the true number on alternating runs."""
    from reports.membership_performance_tracker.logic.receipts_render import (
        coverage_leftovers)
    flags = {"uncovered_billed_accounts": [{"account": "A1", "class": "x"}],
             "billed_excluded_by_rule": [{"account": "A2", "class": "y"}]}
    left = coverage_leftovers({"A1", "A2"}, [], flags)
    assert {r["account"] for r in left} == {"A1", "A2"}
