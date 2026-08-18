#!/usr/bin/env python3
"""tm_formula_audit.py — mechanical audit of every TM-sheet formula.

Born 8/5 after Jiho's gate-1 skim caught two wiring-defect CLASSES the eye
almost missed and tests never covered (formulas live in the template, not
in Python):

  CLASS 1 — mixed-metric SUMIFS: a formula whose SUMIFS terms reference
     metrics that are NOT a Target/Non-Target pair of the same base metric
     (the 8/4 comp rows summed "Retention - Comped ($)" + the ENTIRE
     "Revenue Retained - Non-Target ($)" — a replace-first-occurrence bug).
  CLASS 2 — foreign-row quarter cells: a quarter/summary cell whose
     SUM/AVERAGE arguments reference a DIFFERENT row than the cell's own
     (the 8/5 locations insert copied row-relative formulas verbatim).
  CLASS 3 — naked cross-row addition: A+B over cells that can hold the
     text "-" (unphotographed months) → #VALUE!. Must be N()-wrapped.

Usage:
    /usr/bin/python3 scripts/tm_formula_audit.py <workbook.xlsx>
Exit 1 if any violation. Run against the TEMPLATE and against OUTPUTS.
"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SUMIFS_METRIC = re.compile(r"'Data - Monthly Metrics'!D:D,\"([^\"]+)\"")

# Pure-sum metric pairs that are LEGITIMATE by design (the row's label says
# it sums two families). Everything else mixed in a pure sum is a defect.
LEGIT_SUM_PAIRS = {
    frozenset({"Billables - Allied (#)",
               "Billables - Hospitality Non-Target (#)"}),   # N6 trend rows
}
CELL_REF = re.compile(r"(?<![A-Z:$])(\$?)([A-Z]{1,2})(\$?)(\d+)(?!\()")
AGG = re.compile(r"(?:SUM|AVERAGE)\(([^)]*)\)")
NAKED_ADD = re.compile(r"=\s*([A-Z]{1,2}\d+)\s*\+\s*([A-Z]{1,2}\d+)\s*$")


def base_metric(name: str) -> str:
    return (name.replace(" - Target", " - ").replace(" - Non-Target", " - ")
            .replace("Target ", "").replace("Non-Target ", ""))


def audit(path: Path):
    wb = load_workbook(path)
    violations = []
    for name in wb.sheetnames:
        if not name.startswith("TM -"):
            continue
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=160, max_col=30):
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                where = f"{name}!{cell.coordinate}"
                label = str(ws.cell(cell.row, 1).value or "").strip()

                # CLASS 1 — SUMIFS terms must share one base metric.
                # Ratio formulas (they contain a division) legitimately mix
                # families (trend %s, penetration %s) — the defect class is
                # mixing inside a PURE SUM (8/4 comp rows: comp metric + the
                # entire NT revenue, a replace-first-occurrence artifact).
                metrics = SUMIFS_METRIC.findall(v)
                if (len(set(metrics)) > 1 and "/" not in v
                        and frozenset(metrics) not in LEGIT_SUM_PAIRS):
                    bases = {base_metric(m) for m in metrics}
                    if len(bases) > 1:
                        violations.append(
                            (where, label, "MIXED-METRIC",
                             " + ".join(sorted(set(metrics)))))

                # CLASS 2 — SUM/AVERAGE args on quarter cells reference own row
                for args in AGG.findall(v):
                    refs = CELL_REF.findall(args)
                    ranged = ":" in args
                    if refs and not ranged:
                        rows_ref = {int(r[3]) for r in refs}
                        if rows_ref and rows_ref != {cell.row}:
                            violations.append(
                                (where, label, "FOREIGN-ROW-AGG",
                                 f"args reference row(s) {sorted(rows_ref)}"))

                # CLASS 3 — naked A+B addition (no N() guard)
                m = NAKED_ADD.match(v.replace(" ", ""))
                if m:
                    violations.append(
                        (where, label, "NAKED-ADD", v))
    return violations


def main() -> int:
    path = Path(sys.argv[1])
    violations = audit(path)
    if not violations:
        print(f"✓ {path.name}: every TM-sheet formula passes all 3 classes")
        return 0
    print(f"⛔ {path.name}: {len(violations)} violation(s)")
    seen = set()
    for where, label, kind, detail in violations:
        key = (label, kind, detail.split("+")[0][:30])
        tag = "" if key not in seen else "  (repeat)"
        seen.add(key)
        print(f"  [{kind}] {where}  {label[:40]!r}\n      {detail[:110]}{tag}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
